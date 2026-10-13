from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
GUIDE_PATH = DOCS / "PROJECT_GUIDE.md"
QA_PATH = DOCS / "INTERVIEW_QA.md"
ASSET_DIR = DOCS / "assets"
OUTPUT_PATH = Path(
    os.getenv(
        "FACTORYPILOT_DOCX_OUTPUT",
        str(DOCS / "FactoryPilot项目说明与面试手册.docx"),
    )
)

NAVY = "0B1F33"
NAVY_2 = "153451"
BLUE = "2563EB"
TEAL = "0F766E"
TEAL_LIGHT = "E7F7F4"
BLUE_LIGHT = "EAF2FF"
GRAY_50 = "F8FAFC"
GRAY_100 = "F1F5F9"
GRAY_200 = "E2E8F0"
GRAY_500 = "64748B"
GRAY_700 = "334155"
RED = "B42318"
AMBER = "B45309"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge_data in edges.items():
        edge = tc_borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                edge.set(qn(f"w:{key}"), str(edge_data[key]))


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_table_row_split(row) -> None:
    """Keep a table row on one page so interview answers do not break mid-cell."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_paragraph_border(paragraph, bottom_color=GRAY_200, size="8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), bottom_color)
    p_bdr.append(bottom)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = rgb(color)


def add_hyperlink(paragraph, label: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    run_props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run_props.append(fonts)
    new_run.append(run_props)
    text = OxmlElement("w:t")
    text.text = label
    new_run.append(text)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*]+\*(?!\*))"
)


def add_inline(paragraph, text: str, default_size=10.5, default_color=GRAY_700) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, bold=True, color=NAVY)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(default_size - 0.5, 8.5), color=TEAL)
            run._element.get_or_add_rPr().append(_run_shading(GRAY_100))
        elif token.startswith("["):
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=default_size, color=default_color)
            run.italic = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=default_size, color=default_color)


def _run_shading(fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    return shd


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = rgb("000000")
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(GRAY_700)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for level, size, before, after in [
        (1, 18, 18, 8),
        (2, 14, 14, 6),
        (3, 11.5, 10, 4),
        (4, 10.5, 8, 3),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "FP Code" not in styles:
        code_style = styles.add_style("FP Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["FP Code"]
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = rgb(NAVY)
    code_style.paragraph_format.left_indent = Cm(0.35)
    code_style.paragraph_format.right_indent = Cm(0.35)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(3)
    code_style.paragraph_format.line_spacing = 1.05

    if "FP Quote" not in styles:
        quote_style = styles.add_style("FP Quote", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote_style = styles["FP Quote"]
    quote_style.font.name = "Microsoft YaHei"
    quote_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    quote_style.font.size = Pt(9.5)
    quote_style.font.color.rgb = rgb(GRAY_500)
    quote_style.paragraph_format.left_indent = Cm(0.5)
    quote_style.paragraph_format.right_indent = Cm(0.3)
    quote_style.paragraph_format.space_before = Pt(4)
    quote_style.paragraph_format.space_after = Pt(6)


def configure_page(section, landscape=False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=8.5, color=GRAY_500)


def configure_header_footer(section, first_page=False) -> None:
    section.different_first_page_header_footer = first_page
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.text = ""
    run = p.add_run("FACTORYPILOT  /  MES AGENT ENGINEERING")
    set_run_font(run, size=8, bold=True, color="000000")

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    left = p.add_run("FactoryPilot 项目说明与面试手册    ·    ")
    set_run_font(left, size=8, color=GRAY_500)
    add_page_number(p)

    if first_page:
        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_toc(doc: Document) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run("目录")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(subtitle, "打开 Word 后目录会自动更新；如页码未刷新，选中目录后按 F9。", 9, GRAY_500)

    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "目录将在 Word 中自动生成"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, hint, end])


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AGENT APPLICATION ENGINEERING HANDBOOK")
    set_run_font(r, size=9, bold=True, color=GRAY_500)
    p.paragraph_format.space_after = Pt(22)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FactoryPilot 项目说明与面试手册")
    set_run_font(r, size=30, bold=True, color="000000")
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MES 生产异常诊断与闭环处置 Agent")
    set_run_font(r, size=19, bold=True, color=NAVY_2)
    p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("详细项目说明  ·  技术架构  ·  工程实现  ·  面试问答")
    set_run_font(r, size=11, color=GRAY_500)
    p.paragraph_format.space_after = Pt(28)

    tags = doc.add_table(rows=2, cols=3)
    tags.alignment = WD_TABLE_ALIGNMENT.CENTER
    tags.autofit = False
    for idx, label in enumerate(
        ["LangGraph", "MCP Tools", "Milvus RAG", "Redis Celery", "Neo4j", "A2A"]
    ):
        cell = tags.cell(idx // 3, idx % 3)
        cell.width = Cm(5.0)
        set_cell_shading(cell, TEAL_LIGHT if idx % 2 == 0 else BLUE_LIGHT)
        set_cell_margins(cell, 130, 100, 130, 100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, size=9.5, bold=True, color=TEAL if idx % 2 == 0 else BLUE)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "6", "color": WHITE},
            bottom={"val": "single", "sz": "6", "color": WHITE},
            start={"val": "single", "sz": "6", "color": WHITE},
            end={"val": "single", "sz": "6", "color": WHITE},
        )

    doc.add_paragraph()
    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    values = [
        ("文档版本", "v2.0 分布式执行与多源证据增强版"),
        ("更新日期", date.today().isoformat()),
        ("适用岗位", "Agent 应用开发 / Java 后端 / AI 应用工程"),
        ("数据声明", "全部为确定性合成数据，不包含实习公司真实数据"),
    ]
    for row, (key, value) in zip(info.rows, values):
        row.cells[0].width = Cm(3.2)
        row.cells[1].width = Cm(11.8)
        set_cell_shading(row.cells[0], NAVY)
        set_cell_shading(row.cells[1], GRAY_50)
        for cell in row.cells:
            set_cell_margins(cell, 130, 160, 130, 160)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": GRAY_200},
                bottom={"val": "single", "sz": "4", "color": GRAY_200},
                start={"val": "single", "sz": "4", "color": GRAY_200},
                end={"val": "single", "sz": "4", "color": GRAY_200},
            )
        p1 = row.cells[0].paragraphs[0]
        r1 = p1.add_run(key)
        set_run_font(r1, size=9.5, bold=True, color=WHITE)
        p2 = row.cells[1].paragraphs[0]
        r2 = p2.add_run(value)
        set_run_font(r2, size=9.5, color=GRAY_700)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    r = p.add_run("核心设计：确定性分析提供统计证据，Agent 负责受控推理，人工审批保护高风险执行。")
    set_run_font(r, size=10, bold=True, color=TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("D:\\秋招项目MESAI")
    set_run_font(r, name="Consolas", size=8.5, color=GRAY_500)
    p.paragraph_format.space_before = Pt(30)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_executive_summary(doc: Document) -> None:
    doc.add_heading("阅读指南", level=1)
    p = doc.add_paragraph()
    add_inline(
        p,
        "这份手册由两部分组成：第一部分用于完整理解和展示项目；第二部分用于准备面试追问。建议先掌握端到端闭环，再根据目标岗位重点复习 Java 网关、Agent 编排、安全或评测。",
    )

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cards = [
        ("项目主线", "多源取证 → 统计分析 → RCA → 审批 → 执行 → 验证 → 报告"),
        ("竞争力", "MES 业务理解 + Java 后端 + Agent 工作流 + 工具安全 + 可评测部署"),
        (
            "表达边界",
            "合成数据与单主机部署；Redis/Celery、Neo4j、A2A 已实现，真实厂商服务与共享数据库 checkpoint 仍属后续工作",
        ),
    ]
    for row, (title, body) in zip(table.rows, cards):
        row.cells[0].width = Cm(3.2)
        row.cells[1].width = Cm(12.2)
        set_cell_shading(row.cells[0], TEAL_LIGHT)
        set_cell_shading(row.cells[1], GRAY_50)
        for cell in row.cells:
            set_cell_margins(cell, 140, 160, 140, 160)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": WHITE},
                bottom={"val": "single", "sz": "4", "color": WHITE},
                start={"val": "single", "sz": "4", "color": WHITE},
                end={"val": "single", "sz": "4", "color": WHITE},
            )
        r1 = row.cells[0].paragraphs[0].add_run(title)
        set_run_font(r1, size=9.5, bold=True, color=TEAL)
        add_inline(row.cells[1].paragraphs[0], body, 9.5, GRAY_700)

    doc.add_heading("项目核心数字", level=2)
    metrics = doc.add_table(rows=2, cols=5)
    metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (value, label) in enumerate(
        [
            ("10", "设备"),
            ("500", "生产批次"),
            ("3", "MCP Server"),
            ("14", "容器服务"),
            ("80", "离线评测样例"),
        ]
    ):
        top = metrics.cell(0, idx)
        bottom = metrics.cell(1, idx)
        set_cell_shading(top, NAVY)
        set_cell_shading(bottom, GRAY_50)
        top.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        bottom.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = top.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(value)
        set_run_font(r1, size=18, bold=True, color=WHITE)
        p2 = bottom.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(label)
        set_run_font(r2, size=8.5, bold=True, color=GRAY_500)
        for cell in [top, bottom]:
            set_cell_margins(cell, 110, 80, 110, 80)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "6", "color": WHITE},
                bottom={"val": "single", "sz": "6", "color": WHITE},
                start={"val": "single", "sz": "6", "color": WHITE},
                end={"val": "single", "sz": "6", "color": WHITE},
            )

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_toc(doc)


def find_font(size: int, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def arrow(draw, start, end, color="#64748B", width=5):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 16 * direction, y2 - 9), (x2 - 16 * direction, y2 + 9)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 9, y2 - 16 * direction), (x2 + 9, y2 - 16 * direction)]
    draw.polygon(points, fill=color)


def rounded_box(draw, xy, title, subtitle="", fill="#FFFFFF", outline="#CBD5E1", accent=None):
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    if accent:
        draw.rounded_rectangle((x1, y1, x1 + 12, y2), radius=6, fill=accent)
    title_font = find_font(25, bold=True)
    body_font = find_font(18)
    title_box = draw.textbbox((0, 0), title, font=title_font)
    title_width = title_box[2] - title_box[0]
    draw.text(((x1 + x2 - title_width) / 2, y1 + 24), title, font=title_font, fill="#0B1F33")
    if subtitle:
        subtitle_box = draw.textbbox((0, 0), subtitle, font=body_font)
        subtitle_width = subtitle_box[2] - subtitle_box[0]
        draw.text(((x1 + x2 - subtitle_width) / 2, y1 + 64), subtitle, font=body_font, fill="#64748B")


def create_architecture_diagram(path: Path) -> None:
    image = Image.new("RGB", (1900, 1180), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    title_font = find_font(38, bold=True)
    small_font = find_font(20)
    draw.text((80, 55), "FactoryPilot 总体架构", font=title_font, fill="#0B1F33")
    draw.text(
        (80, 112),
        "网关负责身份与权限，Redis/Celery 负责异步调度，Worker 编排 MCP、图谱和外部 Agent",
        font=small_font,
        fill="#64748B",
    )

    top_y = 205
    boxes = [
        (60, top_y, 330, top_y + 120, "React 控制台", "任务 / 时间线 / 审批", "#E7F7F4", "#0F766E"),
        (390, top_y, 720, top_y + 120, "Spring Boot 网关", "JWT / RBAC / SSE", "#EAF2FF", "#2563EB"),
        (780, top_y, 1070, top_y + 120, "FastAPI API", "创建任务 / 查询状态", "#FFFFFF", "#0F766E"),
        (1130, top_y, 1400, top_y + 120, "Redis", "Broker / 锁 / 通知", "#FFF7ED", "#F59E0B"),
        (1460, top_y, 1790, top_y + 120, "Celery Worker", "LangGraph / 重试 / 超时", "#FFFFFF", "#0F766E"),
    ]
    for x1, y1, x2, y2, title, subtitle, fill, accent in boxes:
        rounded_box(draw, (x1, y1, x2, y2), title, subtitle, fill, "#CBD5E1", accent)
    for start, end in [
        ((330, 265), (390, 265)),
        ((720, 265), (780, 265)),
        ((1070, 265), (1130, 265)),
        ((1400, 265), (1460, 265)),
    ]:
        arrow(draw, start, end, "#2563EB")

    tool_y = 520
    tool_boxes = [
        (65, tool_y, 365, tool_y + 110, "mes-mcp", "只读生产与质量数据", "#EAF2FF", "#2563EB"),
        (425, tool_y, 725, tool_y + 110, "knowledge-mcp", "Milvus 手册与 SOP", "#E7F7F4", "#0F766E"),
        (785, tool_y, 1085, tool_y + 110, "action-mcp", "审批后受控写入", "#FFF7ED", "#F59E0B"),
        (1145, tool_y, 1445, tool_y + 110, "Neo4j", "多跳故障关系证据", "#EAF2FF", "#2563EB"),
        (1505, tool_y, 1805, tool_y + 110, "A2A 厂商 Agent", "Agent Card / JSON-RPC", "#E7F7F4", "#0F766E"),
    ]
    for x1, y1, x2, y2, title, subtitle, fill, accent in tool_boxes:
        rounded_box(draw, (x1, y1, x2, y2), title, subtitle, fill, "#CBD5E1", accent)

    worker_center = 1625
    bus_y = 430
    draw.line((215, bus_y, worker_center, bus_y), fill="#0F766E", width=5)
    arrow(draw, (worker_center, top_y + 120), (worker_center, bus_y), "#0F766E", 5)
    for center in [215, 575, 935, 1295, 1655]:
        arrow(draw, (center, bus_y), (center, tool_y), "#0F766E", 4)

    data_y = 805
    data_boxes = [
        (65, data_y, 365, data_y + 110, "MySQL", "业务状态 / 审批 / 审计"),
        (425, data_y, 725, data_y + 110, "Milvus", "向量召回 / 融合重排"),
        (785, data_y, 1085, data_y + 110, "SQLite Checkpoint", "单主机图状态恢复"),
        (1505, data_y, 1805, data_y + 110, "LangSmith", "Trace / Dataset / Eval"),
    ]
    for x1, y1, x2, y2, title, subtitle in data_boxes:
        rounded_box(draw, (x1, y1, x2, y2), title, subtitle, "#FFFFFF", "#CBD5E1", "#64748B")

    arrow(draw, (215, tool_y + 110), (215, data_y), "#64748B", 4)
    arrow(draw, (575, tool_y + 110), (575, data_y), "#64748B", 4)
    draw.line((935, tool_y + 110, 935, 735), fill="#64748B", width=4)
    draw.line((215, 735, 935, 735), fill="#64748B", width=4)
    arrow(draw, (215, 735), (215, data_y), "#64748B", 4)
    arrow(draw, (935, 735), (935, data_y), "#64748B", 4)
    arrow(draw, (1655, tool_y + 110), (1655, data_y), "#64748B", 4)

    footer_font = find_font(18)
    draw.text((80, 1040), "MySQL 保存业务事实；Redis 用于协调和临时状态；SQLite checkpoint 当前只用于单主机恢复。", font=footer_font, fill="#334155")
    draw.text((80, 1080), "写操作保护：JWT/RBAC → interrupt 审批 → Action MCP 二次查验 → 幂等与审计", font=footer_font, fill="#0F766E")
    image.save(path, quality=95)


def create_workflow_diagram(path: Path) -> None:
    image = Image.new("RGB", (1900, 1220), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    title_font = find_font(38, bold=True)
    body_font = find_font(19)
    draw.text((80, 48), "LangGraph 异常处置工作流", font=title_font, fill="#0B1F33")
    draw.text((80, 103), "确定性节点完成计算与安全控制，RCA 节点在证据范围内使用模型推理", font=body_font, fill="#64748B")

    top = [
        (80, 205, 310, 315, "用户异常", "自然语言输入"),
        (390, 205, 620, 315, "Triage", "分类 / 严重度"),
    ]
    for box in top:
        rounded_box(draw, box[:4], box[4], box[5], "#FFFFFF", "#CBD5E1", "#2563EB")
    arrow(draw, (310, 260), (390, 260), "#2563EB")

    collectors = [
        (760, 175, 990, 275, "质量与批次", "mes-mcp"),
        (1040, 175, 1270, 275, "设备告警", "mes-mcp"),
        (760, 325, 990, 425, "维修记录", "mes-mcp"),
        (1040, 325, 1270, 425, "SOP / 手册", "knowledge-mcp"),
    ]
    for box in collectors:
        rounded_box(draw, box[:4], box[4], box[5], "#EAF2FF", "#BFDBFE", "#2563EB")
    arrow(draw, (620, 260), (720, 260), "#2563EB")
    draw.line((720, 225, 720, 375), fill="#2563EB", width=5)
    for y in [225, 375]:
        arrow(draw, (720, y), (760, y), "#2563EB", 4)
        arrow(draw, (720, y), (1040, y), "#2563EB", 4)

    draw.line((1320, 225, 1320, 375), fill="#0F766E", width=5)
    for x, y in [(990, 225), (1270, 225), (990, 375), (1270, 375)]:
        arrow(draw, (x, y), (1320, y), "#0F766E", 4)

    rounded_box(draw, (1400, 245, 1700, 355), "确定性统计分析", "基线 / 相关性 / 重合度", "#E7F7F4", "#99E4D7", "#0F766E")
    arrow(draw, (1320, 300), (1400, 300), "#0F766E")

    rounded_box(draw, (1180, 445, 1430, 545), "Neo4j 图谱", "故障多跳路径", "#EAF2FF", "#BFDBFE", "#2563EB")
    rounded_box(draw, (1500, 445, 1750, 545), "A2A 厂商 Agent", "外部专家意见", "#E7F7F4", "#99E4D7", "#0F766E")
    draw.line((1550, 355, 1550, 405), fill="#0F766E", width=5)
    draw.line((1305, 405, 1625, 405), fill="#0F766E", width=5)
    arrow(draw, (1305, 405), (1305, 445), "#0F766E", 4)
    arrow(draw, (1625, 405), (1625, 445), "#0F766E", 4)

    second_row = [
        (1400, 665, 1700, 775, "RCA Agent", "汇总全部证据", "#FFFFFF", "#0F766E"),
        (1030, 665, 1320, 775, "处置方案", "风险分级 / 参数约束", "#FFFFFF", "#0F766E"),
        (660, 665, 950, 775, "人工审批", "LangGraph interrupt", "#FFF7ED", "#F59E0B"),
        (290, 665, 580, 775, "Action MCP", "创建工单 / 调整抽检", "#FFFFFF", "#2563EB"),
    ]
    for box in second_row:
        rounded_box(draw, box[:4], box[4], box[5], box[6], "#CBD5E1", box[7])
    arrow(draw, (1305, 545), (1450, 665), "#0F766E", 4)
    arrow(draw, (1625, 545), (1625, 665), "#0F766E", 4)
    arrow(draw, (1400, 720), (1320, 720), "#0F766E")
    arrow(draw, (1030, 720), (950, 720), "#F59E0B")
    arrow(draw, (660, 720), (580, 720), "#2563EB")

    rounded_box(draw, (290, 930, 580, 1040), "结果验证", "连续 3 批恢复条件", "#E7F7F4", "#99E4D7", "#0F766E")
    rounded_box(draw, (660, 930, 950, 1040), "事故报告", "证据 / 处置 / 结论", "#EAF2FF", "#BFDBFE", "#2563EB")
    arrow(draw, (435, 775), (435, 930), "#0F766E")
    arrow(draw, (580, 985), (660, 985), "#0F766E")

    label_font = find_font(17, bold=True)
    draw.text((590, 680), "批准", font=label_font, fill="#2563EB")
    draw.text((735, 805), "拒绝 → 记录原因并安全结束", font=body_font, fill="#B45309")
    draw.text((80, 1135), "RCA 合并结构化数据、统计、SOP、图谱和厂商意见；写操作必须在审批后执行。", font=body_font, fill="#0F766E")
    image.save(path, quality=95)


def split_table_row(line: str):
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in re.split(r"(?<!\\)\|", stripped)]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows = [rows[0]] + rows[2:]
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_table_row_split(row)
    for row_idx, source_row in enumerate(rows):
        for col_idx in range(columns):
            cell = table.cell(row_idx, col_idx)
            value = source_row[col_idx] if col_idx < len(source_row) else ""
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 95, 105, 95, 105)
            if row_idx == 0:
                set_cell_shading(cell, NAVY)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, GRAY_50)
            else:
                set_cell_shading(cell, WHITE)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "4", "color": GRAY_200},
                bottom={"val": "single", "sz": "4", "color": GRAY_200},
                start={"val": "single", "sz": "4", "color": GRAY_200},
                end={"val": "single", "sz": "4", "color": GRAY_200},
            )
            p = cell.paragraphs[0]
            add_inline(p, value, 8.5 if columns > 3 else 9, WHITE if row_idx == 0 else GRAY_700)
            for run in p.runs:
                if row_idx == 0:
                    run.font.color.rgb = rgb(WHITE)
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, code: str, language: str) -> None:
    p = doc.add_paragraph(style="FP Code")
    shade_paragraph(p, GRAY_100)
    p.paragraph_format.keep_together = True
    if language:
        label = p.add_run(f"{language.upper()}\n")
        set_run_font(label, name="Consolas", size=7.5, bold=True, color=TEAL)
    run = p.add_run(code.rstrip())
    set_run_font(run, name="Consolas", size=8.5, color=NAVY)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), TEAL)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def render_markdown(doc: Document, markdown: str, mermaid_images: list[Path], heading_shift=-1) -> None:
    lines = markdown.replace("\r\n", "\n").split("\n")
    index = 0
    mermaid_index = 0
    skip_first_h1 = True
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        fence_match = re.match(r"^```([^`]*)$", stripped)
        if fence_match:
            language = fence_match.group(1).strip()
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if language == "mermaid" and mermaid_index < len(mermaid_images):
                image_path = mermaid_images[mermaid_index]
                mermaid_index += 1
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                picture = run.add_picture(str(image_path), width=Cm(16.4))
                alt_text = "FactoryPilot 系统总体架构图" if mermaid_index == 1 else "FactoryPilot LangGraph 异常处置工作流图"
                picture._inline.docPr.set("descr", alt_text)
                picture._inline.docPr.set("title", alt_text)
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = caption.add_run(f"图 {mermaid_index}  ·  " + ("系统总体架构" if mermaid_index == 1 else "LangGraph 工作流"))
                set_run_font(cap_run, size=8.5, color=GRAY_500)
            else:
                add_code_block(doc, "\n".join(code_lines), language)
            index += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading_match:
            markdown_level = len(heading_match.group(1))
            title = heading_match.group(2).strip()
            if markdown_level == 1 and skip_first_h1:
                skip_first_h1 = False
                index += 1
                continue
            word_level = max(1, min(4, markdown_level + heading_shift))
            p = doc.add_paragraph(style=f"Heading {word_level}")
            if title.startswith("14. 不建议在面试中使用的表述"):
                p.paragraph_format.page_break_before = True
            add_inline(p, title, [0, 18, 14, 11.5, 10.5][word_level], NAVY)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(split_table_row(lines[index]))
                index += 1
            add_markdown_table(doc, table_lines)
            continue

        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            index += 1
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            p = doc.add_paragraph(style="FP Quote")
            add_inline(p, "\n".join(quote_lines), 9.5, GRAY_500)
            continue

        bullet_match = re.match(r"^\s*[-*]\s+(.+)$", line)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.25)
            p.paragraph_format.space_after = Pt(2.5)
            add_inline(p, bullet_match.group(1))
            index += 1
            continue

        number_match = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if number_match:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.25)
            p.paragraph_format.space_after = Pt(2.5)
            add_inline(p, number_match.group(1))
            index += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        index += 1


def add_part_page(doc: Document, number: str, title: str, subtitle: str, color: str) -> None:
    doc.add_page_break()
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"PART {number}")
    set_run_font(r, size=11, bold=True, color=color)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, size=26, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_run_font(r, size=11, color=GRAY_500)
    p.paragraph_format.space_after = Pt(30)
    doc.add_page_break()


def add_quick_review(doc: Document) -> None:
    add_part_page(doc, "03", "面试前快速复习卡", "用 5 分钟检查项目主线、边界和演示流程", AMBER)
    doc.add_heading("一句话定位", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    add_inline(
        p,
        "FactoryPilot 是一个通过确定性统计证据、LangGraph 有状态编排和人工审批，实现 MES 生产异常从诊断到验证闭环的受控 Agent 系统。",
        11,
        NAVY,
    )

    doc.add_heading("必须记住的 10 个关键词", level=1)
    keywords = [
        ("确定性分析", "LLM 不直接猜统计结论"),
        ("证据 ID", "根因必须引用真实证据"),
        ("并行采集", "质量、告警、维修、SOP"),
        ("图谱证据", "Neo4j 返回故障多跳路径"),
        ("厂商协作", "A2A 只连接独立外部 Agent"),
        ("interrupt", "高风险动作暂停等待审批"),
        ("双重授权", "网关 RBAC + Action MCP 查库"),
        ("幂等审计", "重复执行不重复创建工单"),
        ("Redis Celery", "跨进程锁与可重试任务队列"),
        ("真实边界", "合成数据与单主机 checkpoint"),
    ]
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (title, body) in enumerate(keywords):
        cell = table.cell(idx // 2, idx % 2)
        set_cell_shading(cell, GRAY_50 if idx % 2 else BLUE_LIGHT)
        set_cell_margins(cell, 130, 150, 130, 150)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "6", "color": WHITE},
            bottom={"val": "single", "sz": "6", "color": WHITE},
            start={"val": "single", "sz": "6", "color": WHITE},
            end={"val": "single", "sz": "6", "color": WHITE},
        )
        p = cell.paragraphs[0]
        r = p.add_run(title + "\n")
        set_run_font(r, size=10, bold=True, color=TEAL)
        r = p.add_run(body)
        set_run_font(r, size=8.5, color=GRAY_500)

    doc.add_heading("现场演示 9 步", level=1)
    for text in [
        "使用 supervisor 账号登录，展示 JWT/RBAC。",
        "输入 2 号产线激光打码良率异常，启动诊断。",
        "展示任务由 Redis Broker 投递并由 Celery Worker 执行。",
        "展示四路数据采集和 LangGraph 时间线。",
        "查看最近 20 批趋势、均值差、相关性和告警重合证据。",
        "查看 Neo4j 故障路径和 A2A 厂商专家意见。",
        "展示带引用的根因假设和风险处置计划。",
        "在等待审批处暂停，批准后创建幂等维修工单。",
        "查看恢复验证、事故报告、审计日志和 LangSmith Trace。",
    ]:
        p = doc.add_paragraph(style="List Number")
        add_inline(p, text)

    doc.add_heading("绝对不要说错的边界", level=1)
    boundaries = [
        "Redis/Celery、Neo4j 和 A2A 已实现并完成 Docker 端到端验证；Kubernetes 未引入。",
        "HashEmbedding 是离线演示方案，不是生产级中文 embedding。",
        "SQLite checkpoint 只适合当前单主机共享卷；多主机部署需要数据库型共享 checkpointer。",
        "Redis run lease 尚未自动续租，因此任务硬超时必须小于锁 TTL。",
        "当前 A2A 厂商 Agent 和 Neo4j 故障知识均为可复现的演示实现，不代表已连接真实厂商系统。",
        "评测基于合成数据；局部 100% 不等于生产准确率。",
        "Agent 生成根因假设，不保证自动找到唯一真实根因。",
    ]
    for text in boundaries:
        p = doc.add_paragraph(style="List Bullet")
        add_inline(p, text)


def build_document() -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    architecture_path = ASSET_DIR / "word_architecture.png"
    workflow_path = ASSET_DIR / "word_workflow.png"
    create_architecture_diagram(architecture_path)
    create_workflow_diagram(workflow_path)

    doc = Document()
    configure_styles(doc)
    configure_page(doc.sections[0])
    configure_header_footer(doc.sections[0], first_page=True)
    set_update_fields(doc)

    core_props = doc.core_properties
    core_props.title = "FactoryPilot MES 生产异常诊断与闭环处置项目说明与面试手册"
    core_props.subject = "MES 生产异常诊断与闭环处置 Agent"
    core_props.author = "FactoryPilot Project"
    core_props.keywords = "MES, Agent, LangGraph, MCP, RAG, Milvus, Redis, Celery, Neo4j, A2A, Spring Boot, FastAPI"
    core_props.comments = "由仓库实际实现和项目文档整合生成"

    add_cover(doc)
    add_executive_summary(doc)

    add_part_page(
        doc,
        "01",
        "详细项目说明",
        "从业务问题、系统架构到安全、评测与部署的完整技术手册",
        TEAL,
    )
    render_markdown(
        doc,
        GUIDE_PATH.read_text(encoding="utf-8"),
        [architecture_path, workflow_path],
        heading_shift=-1,
    )

    add_part_page(
        doc,
        "02",
        "面试问题与参考答案",
        "105 道问题，覆盖项目介绍、架构、Agent、数据分析、RAG、安全与生产化",
        BLUE,
    )
    render_markdown(doc, QA_PATH.read_text(encoding="utf-8"), [], heading_shift=-1)
    add_quick_review(doc)

    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section, first_page=section is doc.sections[0])

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_document()
    print(output)
