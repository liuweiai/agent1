"""Reader-facing source walkthrough inserted into the interview handbook."""
from pathlib import Path
from docx.shared import Pt
from walkthrough_figures import mcp_figure

ROOT = Path(__file__).resolve().parents[1]
PY = 'agent-service/src/factorypilot/'
JAVA = 'mes-gateway/src/main/java/com/factorypilot/gateway/'


def add_details(doc):
    from build_code_walkthrough_handbook import add_paragraph as p, add_code as code, add_table as table, add_figure as figure

    def h(text, level=2):
        doc.add_heading(text, level=level)

    def source(path, marker, count=0):
        lines = (ROOT/path).read_text(encoding='utf-8').splitlines()
        i = next(i for i, line in enumerate(lines) if marker in line)
        q = doc.add_paragraph(f'源码定位  {path}:{i+1}')
        q.paragraph_format.space_after = Pt(3)
        q.paragraph_format.keep_with_next = True
        for r in q.runs:
            r.font.size = Pt(9)
        if count:
            selected = lines[i:i+count]
            indent = min((len(s)-len(s.lstrip()) for s in selected if s.strip()), default=0)
            code(doc, '\n'.join(f'{i+j+1:>3}  {s[indent:]}' for j,s in enumerate(selected)))

    h('源码逐调用精读', 1)
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    p(doc, '这一章使用“2号产线激光打码良率从98%降到91%”贯穿讲解。用 R 表示本次随机生成的 run_id，避免把示例 ID 当成真实运行记录。每读一节，在编辑器打开给出的文件，搜索函数名，核对它调用了谁、拿了什么数据、返回了什么字段。')
    h('入口一  前端按钮和 API 请求')
    source('frontend/src/App.tsx', 'await api.createRun', 5)
    source('frontend/src/api.ts', 'createRun:', 3)
    p(doc, '按钮触发创建逻辑，query.trim() 去掉前后空白，api.createRun 使用默认 LINE-02。API 辅助函数序列化 query 与 line_id，fetch 使用 credentials: include 携带 Cookie。当前产线参数来自请求字段，不能把“自然语言识别任意产线”说成已经实现。前端得到任务 ID 后选择该任务，订阅事件并读取详情。')
    h('入口二  网关认证与代理')
    source(JAVA+'config/SecurityConfig.java', '.hasAnyRole', 1)
    source(JAVA+'web/AgentProxyController.java', 'public Mono<ResponseEntity<String>> createRun', 4)
    p(doc, 'Spring Security 先验证身份：普通业务接口要求已登录，审批接口额外要求 SUPERVISOR 或 ADMIN。createRun 控制器使用 WebClient 把请求转发到 Agent Service。追 WebClient 配置可找到 x-gateway-secret；FastAPI 在 REQUIRE_GATEWAY_AUTH=true 时用 hmac.compare_digest 校验它。这是服务之间的信任校验，用户 JWT 和服务共享密钥用途不同。')
    source(PY+'main.py', 'async def verify_gateway', 7)
    h('入口三  创建任务和投递队列')
    source(PY+'main.py', 'def create_run(', 9)
    source(PY+'service.py', 'def create_run(', 4)
    p(doc, 'FastAPI 调 AgentService.create_run，生成 UUID 的 hex 字符串 R，写 runs，并保存 request 事件。随后 enqueue_start 返回调度 task_id，HTTP 返回 202。202 只表示已受理；分析尚未结束，甚至可能还没有 Worker 领取。State 也还没有完整业务数据。')
    source(PY+'task_queue.py', 'def enqueue_start(', 12)
    p(doc, 'TASK_QUEUE_BACKEND=celery 时，send_task 发 factorypilot.run.start，参数只含 R，任务 ID 是 run-start-R。Worker 收到后在 start_run_task 中取得本进程的 AgentService，调用 start_run(R, propagate_transient=True)。本地模式则把 start_run 加入 FastAPI BackgroundTasks，没有跨进程消息。Celery task_id 标识一次调度，run_id 标识整个业务任务，两者不能混用。')
    h('入口四  服务启动时如何装配依赖')
    source(PY+'service.py', 'if settings.tool_transport', 6)
    p(doc, 'AgentService.__init__ 先创建业务数据库和 RedisCoordinator，再创建 SQLite SqliteSaver。AGENT_MODE 决定 DemoRcaEngine 或 OpenAIRcaEngine；TOOL_TRANSPORT 决定 Remote 工具适配器或本地工具对象。构造 FactoryPilotGraph 时把这些对象注入 self.mes、self.knowledge、self.actions。看到 graph 中同一个方法名，并不能确定有没有联网，必须先看注入对象。')
    p(doc, 'Neo4j 和 A2A 客户端另外构造，它们不经过 knowledge-mcp。业务数据库可以选 MySQL，但 checkpointer 仍是 SQLite 文件。开关启用只说明选中了代码路径，实际连接和任务结果还要看运行日志。')
    h('入口五  Worker 进入 LangGraph')
    source(PY+'service.py', 'initial = {', 12)
    source(PY+'service.py', 'def config(run_id', 2)
    p(doc, 'start_run 先通过 _claim 申请进程内保护和可选 Redis 锁，再读 runs。第一次运行没有 checkpoint，因此 invoke(initial)。如果已有 values 且有 next，则 invoke(None) 尝试继续已有任务；审批恢复则有独立的 resume_run 分支。配置将 thread_id 设为 R，recursion_limit=30 是图运行的步数保护，不表示业务会循环30次。')
    code(doc, '本次初始 State（示例）\n{\n  "run_id": "R",\n  "query": "2号产线激光打码设备良率从98%降到91%…",\n  "line_id": "LINE-02",\n  "evidence": []\n}')

    h('State 怎样随着节点变化', 1)
    source(PY+'graph.py', 'class AgentState', 23)
    p(doc, 'AgentState 是 total=False 的 TypedDict，一开始可以只提供部分字段。节点返回字典是状态增量。普通字段按新值更新；evidence 使用 Annotated[list, operator.add]，把各分支的新证据列表相加。它没有自动去重能力，所以节点应该只返回新增证据，不能把旧列表再带回。status、error 和当前步骤不是 AgentState 的字段，它们由业务记录和事件承担。')
    code(doc, '四路分支返回增量的示意（数值与ID仅为说明）\nmes_data         → {device_context: {...}, quality: [...], evidence: [E1]}\nalarm_data       → {alarms: [...], evidence: [E2]}\nmaintenance_data → {maintenance: [...], evidence: [E3]}\nknowledge_data   → {documents: [...], evidence: [E4]}\n汇合后的 evidence = [E1, E2, E3, E4]\n其余四类结果分别进入各自字段，不互相覆盖')
    p(doc, 'statistical_analysis 再写 production_analysis 和三条 STAT-* 证据；Neo4j/A2A 成功时继续追加证据；rca 写 hypotheses；plan 写 action_plan 和 risk_level；approval 恢复后写 approval；execute 写 action_results；verify 写 verification 并更新 production_analysis；report 写 report。State 传递顺序正是你纸上标注输入输出的顺序。')

    h('Edge 怎样保证先后顺序', 1)
    source(PY+'graph.py', 'builder.add_edge(START', 17)
    p(doc, 'for 循环从 triage 连出四条边，因此四个采集节点可以并行。add_edge 的起点是四元素列表时，含义是等待四个节点都完成，再执行统计分析；不能简化成任一路结束就开始分析。同样，RCA 等待 graph_knowledge 和 vendor_expert 两路都结束。禁用的补充节点返回空结果，也算该节点完成。')
    source(PY+'graph.py', 'return Command(update=', 1)
    p(doc, 'approval 的分支由 Command 决定，build 中没有 approval→execute 的静态 add_edge，也没有 add_conditional_edges。纸上需要画出两条条件出口：同意到 execute→verify→report，拒绝到 reject→report。两条路径共用 report，但拒绝不经过 verify。源码没有 verify→plan 回边，也没有“不断改参数直到恢复”的业务循环。')

    h('十五个 Node 逐个跟踪', 1)
    nodes = [
      ('01 triage','def triage(',
       '读取 query 和 line_id，调用 triage.py 的 classify_anomaly 关键词规则。把 line_id 末两位拼为 DEV-L02-LASER，并更新 runs 的运行状态和严重程度。返回 anomaly_type、severity、device_id。它不调用 LLM；产线到设备的拼接也不是动态设备选择算法。',
       '下一步是四路采集同时具备运行条件，四路共享 run_id 以关联审计与事件。'),
      ('02 mes_data','def collect_mes_data(',
       '依次调用 get_line_device(R, LINE-02) 与 get_recent_quality(R, LINE-02, 20)。前者读设备台账；后者 JOIN batches、inspections、process_metrics 后取最近记录。节点把返回结果整理成设备和质量证据，写 device_context、quality、evidence。',
       '同一节点内的两次调用是顺序执行；它与其他三个采集节点之间才是并行。未来一批多检验记录时，JOIN后的20条不一定等于20个不同批次。'),
      ('03 alarm_data','def collect_alarm_data(',
       '使用 triage 派生的 device_id，调用 get_recent_alerts(R, DEV-L02-LASER, 24, 20)。结果保存为 alarms；最多前6条包装成 evidence，提供告警码、发生时间和说明。',
       '原始告警列表用于统计时间重合；证据列表用于RCA引用。两者并非同一个字段。'),
      ('04 maintenance_data','def collect_maintenance_data(',
       '调用 get_maintenance_history(R, DEV-L02-LASER, 10)，按完成时间倒序取维修记录。保存 maintenance，并将前5条转为 evidence，供RCA了解历史故障与处置。',
       '历史维修只能辅助提出假设，不能证明当前设备故障已被修复。'),
      ('05 knowledge_data','def collect_knowledge_data(',
       '把用户query加上固定“光路 镜片 审批”词，调用 knowledge.search(R, query, 4)。远程时MCP函数名是 search_knowledge；服务端可走 lexical 或 Milvus。返回 documents 和带文档ID、段落、版本的 evidence。',
       '固定补词体现当前演示偏向激光场景；检索为空仍能继续，当前缺少显式citation degraded标记，这就是SCN-016失败的原因。'),
      ('06 statistical_analysis','def analyze_production_data(',
       '四路汇合后调用 analyze_production_window(quality, alarms)。按97%阈值区分稳定组与异常组，算均值、百分点变化、速度/功率/温度与缺陷率的Pearson相关性，以及±6小时告警重合。返回 production_analysis，并追加 STAT-WINDOW-1、STAT-CORR-1、STAT-TIME-1。',
       '这是纯Python计算。相关性不是因果；样本不足会标记证据不足，但当前plan并不会据此自动改成补采计划。'),
      ('07 graph_knowledge','def collect_graph_knowledge(',
       '从 device_context 取型号，空时回退LM-5000，调用 search_fault_paths(R, model, query, 4)。Neo4j驱动执行固定参数化关系遍历，返回故障、部件、症状、处置与SOP路径；保存 graph_context 并追加 GRAPH-* 证据。',
       '这是Agent Service直连Neo4j，不是knowledge-mcp调用。配置关闭返回空数组并标记skipped。'),
      ('08 vendor_expert','def consult_vendor_expert(',
       '构造 query、device_id、model、前10个alarm_codes和production_analysis，交 VendorExpertClient.consult。客户端发现Agent Card后，通过JSON-RPC message/send请求独立厂商服务；读取completed任务的data artifact，保存 vendor_expert 和 A2A-VENDOR-1 证据。',
       '厂商服务diagnose当前也是规则函数；没有调用另一个LLM，也没有写生产设备的权限。该客户端不轮询tasks/get。'),
      ('09 rca','def root_cause_analysis(',
       '两路补充全部结束后，调用 rca_engine.analyze(state)。Demo按规则生成2或3个假设；LLM模式把query、统计结果和evidence交给ChatOpenAI结构化输出。返回hypotheses。有效引用ID会被保留，无效ID会被过滤；没有有效引用时置信度最多0.35。',
       '只有这里可以调用主LLM。它没有bind_tools，不负责自由挑选工具；“最多3个假设”只在prompt中，Schema未对列表长度设max_length。'),
      ('10 plan','def create_action_plan(',
       '读取device_id和line_id，构造固定的维修工单动作与抽检加严动作，设置risk_level=high并保存业务风险。输出action_plan和risk_level。',
       '当前实现不读hypotheses，也不让模型生成可执行参数。面试说“固定受控计划”比“任意异常自主规划”准确。'),
      ('11 approval','def request_approval(',
       '先将runs.status改为awaiting_approval并发送waiting事件，再interrupt返回run_id、风险和actions。首次运行在这里暂停；收到Command(resume=approval)后该函数从头重入，interrupt取回审批决定，补decided_at并写runs.approval_json。',
       '返回Command同时更新State.approval并跳execute/reject。等待审批不是HTTP请求一直阻塞，也不是Worker一直睡眠。'),
      ('12 execute','def execute_actions(',
       '依次调用create_maintenance_order和tighten_inspection，创建P1工单并把抽检周期设为10分钟。保存action_results。远程模式先经过Action MCP，服务端再从业务库取批准记录。',
       'execute不遍历action_plan；动作和参数由Python硬编码。工单标题还固定写2号线，因此不能说其他产线都已有完整适配。'),
      ('13 reject','def reject_actions(',
       '发送跳过执行事件并返回action_results=[]。没有创建工单和加严抽检调用。事件step用execute，所以时间线上看见execute/skipped并不代表发生了写入。',
       '它的下一跳是report，拒绝分支不运行verify。'),
      ('14 verify','def verify_result(',
       '直接self.database.fetch_one读取work_orders和line_controls，检查存在工单且抽检周期为10；再以固定[96.8,97.4,97.9]模拟批次调用evaluate_recovery。保存verification并回填production_analysis.recovery_validation。',
       '没有通过MES MCP重新采集生产质量，也没有确认工单已被维修人员完成。末尾只有连续2批≥97%，因此observing；workflow completed不等于质量恢复。'),
      ('15 report','def generate_report(',
       'Python组装report_id、incident、scope、首个hypothesis、证据数、approval、actions、verification和conclusion，写runs.report_json并返回report。',
       '该节点没有第二次LLM总结。图到END后，service先同步state_json，再把runs.status置为completed或rejected，供前端刷新。'),
    ]
    for title, marker, action, boundary in nodes:
        h(title)
        source(PY+'graph.py', marker)
        p(doc, action)
        p(doc, boundary)

    h('一条真实 MCP 调用穿过哪些函数', 1)
    path=ROOT/'docs/assets/mcp_call_walkthrough.png'
    mcp_figure(path)
    figure(doc,path,'图 2  质量查询的进程边界与调用返回路径',16.8,'质量查询从LangGraph节点经过Remote适配器、MCP客户端、服务端工具到数据库')
    p(doc,'图中蓝色中间层仅在 TOOL_TRANSPORT=mcp 时存在。服务端 @mcp.tool 暴露函数，客户端 call_tool 使用相同工具名。四路采集并行并不意味着把三个Server的工具交给模型自由选择。')
    h('公共 RemoteCaller 的入口与结果解包')
    source(PY+'mcp_client_tools.py','async def _call_async',10)
    p(doc,'Remote适配器把业务参数组成arguments交给MCPRemoteCaller.call。call按max_attempts循环，每次asyncio.run(_call_async)。HTTP客户端建立Streamable HTTP传输，MCP Client在上下文内发call_tool。结果is_error时抛MCPToolCallError；否则优先读structured_content，必要时从text解析JSON。解包后的Python列表或字典回到Node，才被包装成State和证据。')
    p(doc,'默认HTTP与MCP读超时配置为8秒，call_tool还被9秒的asyncio.wait_for包裹；连接上下文的建立和退出不在这个wait_for内，不能把8秒说成整个工作流的总截止时间。最多2次尝试，中间退避0.15秒。当前捕获所有Exception重试，参数错误也可能重试。tools.py里的ToolPolicy只是保存配置，本地_read_with_audit没有使用它实施超时或重试；不能把远程保证套到所有本地函数上。')
    h('七种 MCP 工具逐一定位')
    tools=[
      ('get_line_device','mes_data','RemoteMesTools.get_line_device','mes_server.py','def get_line_device','MesReadTools.get_line_device','devices，固定line_id参数SELECT'),
      ('get_recent_quality','mes_data','RemoteMesTools.get_recent_quality','mes_server.py','def get_recent_quality','MesReadTools.get_recent_quality','batches JOIN inspections LEFT JOIN process_metrics'),
      ('get_recent_alerts','alarm_data','RemoteMesTools.get_recent_alerts','mes_server.py','def get_recent_alerts','MesReadTools.get_recent_alerts','alerts，device_id/时间窗口/limit参数'),
      ('get_maintenance_history','maintenance_data','RemoteMesTools.get_maintenance_history','mes_server.py','def get_maintenance_history','MesReadTools.get_maintenance_history','maintenance_records，device_id参数与时间排序'),
      ('search_knowledge','knowledge_data','RemoteKnowledgeTools.search','knowledge_server.py','def search_knowledge','KnowledgeTools.search 或 MilvusKnowledgeTools.search','knowledge_docs；Milvus模式额外向量召回与固定权重融合'),
      ('create_maintenance_order','execute','RemoteActionTools.create_maintenance_order','action_server.py','def create_maintenance_order','SafeActionTools.create_maintenance_order','查runs审批，再查work_orders；不存在则INSERT'),
      ('tighten_inspection','execute','RemoteActionTools.tighten_inspection','action_server.py','def tighten_inspection','SafeActionTools.tighten_inspection','查runs审批与成功audit_logs，再UPDATE line_controls'),
    ]
    for name,node,adapter,server,marker,impl,backend in tools:
        h(name,3)
        p(doc,f'{node} → {adapter} → MCPRemoteCaller.call → client.call_tool("{name}", arguments) → {server}中的{name} → {impl}。')
        method = 'search' if name == 'search_knowledge' else name
        source(PY+'mcp_client_tools.py','def '+method+'(')
        source(PY+'mcp_servers/'+server,marker)
        source(PY+'tools.py','def '+method+'(')
        if name == 'search_knowledge':
            source(PY+'milvus_knowledge.py','def search(')
        p(doc,'后端访问：'+backend+'。')
    p(doc,'图中的计划文案使用action-mcp.create_maintenance_order等带前缀名称，实际call_tool只传create_maintenance_order。审计operation可能写hybrid_search或milvus_hybrid_search，协议工具名仍是search_knowledge。协议名称、Python方法名、审计标签要分开记。')
    h('写工具拿到的审批从哪里来')
    source(PY+'mcp_servers/action_server.py','def _server_verified_approval',6)
    p(doc,'RemoteActionTools的方法签名接收approval，但上传arguments时没有带这个对象。服务端根据run_id重新读取runs.approval，检查存在、approved为真、approver非空，再交SafeActionTools校验。因此MCP写服务必须连到能看到同一个runs记录的业务库。')
    p(doc,'当前服务端没有再次查询用户角色、runs.status或逐项匹配审批计划参数。角色由Spring网关限制；本地模式直接把State.approval交SafeActionTools。当前措施能演示多层受控执行，但不能说已经实现任意部署边界下完整授权或数据库原子exactly-once。')
    h('工具结果和审计如何闭环')
    source(PY+'tools.py','def _read_with_audit',5)
    p(doc,'MES查询进入_read_with_audit后，finally记录actor、tool、operation、parameters、outcome和duration_ms。数据库参数由业务函数绑定，模型没有任意SQL执行入口。参数正则校验发生在这个finally之前，因此某些参数拒绝不会产生该工具审计。知识检索和写动作有自己的审计路径；verify直读数据库也不经过MES工具审计。')
    p(doc,'工单先按run_id+device_id查旧单，已有就返回旧ID；抽检先查本run成功审计，已有就返回幂等结果。它们是先查后写流程，缺少业务唯一约束和原子事务保护。网络重试可能发生在写成功但响应丢失之后，这正是需要幂等的原因，也解释了当前实现仍有并发改进空间。')

    h('审批恢复的时间顺序', 1)
    source(JAVA+'web/AgentProxyController.java','public Mono<ResponseEntity<String>> approve',9)
    source(PY+'service.py','def approve(',15)
    p(doc,'第一步：浏览器POST审批，Java读取已认证身份并覆盖body.approver。第二步：FastAPI调用service.approve，读取awaiting_approval状态，生成approval_id并把状态先改成executing/running，接着投递resume任务。这一刻approval_json尚未写入。检查与UPDATE不是数据库CAS，只能用来解释串行重复审批被拒绝。')
    source(PY+'service.py','self.graph.invoke(Command(resume=',7)
    p(doc,'第三步：Worker调用resume_run并使用相同R作为thread_id。它先记审批事件与审计，再invoke(Command(resume=approval))。第四步：LangGraph重新进入request_approval函数，interrupt之前的状态更新和事件发送再次执行，interrupt这次返回决定；节点写入approval_json，然后Command.goto控制下一跳。')
    p(doc,'第五步：批准到execute，Action MCP此时才读得到持久化批准；拒绝到reject。执行链到END后，_save_graph_state把checkpoint.values序列化进runs.state_json，随后服务层写completed/rejected。checkpoint保留图的恢复点；runs给业务API查询；它们不是一个原子跨服务事务。')
    p(doc,'waiting事件通过(run_id,idempotency_key)唯一约束去重，避免节点重入造成同一等待事件重复插入。这个事件唯一约束不能推导成工单也有唯一约束。Celery有限重试用于技术异常，不能把它讲成良率不恢复就自动重新规划。')

    h('前端时间线如何看到这些变化',1)
    source('frontend/src/App.tsx','const source = new EventSource',10)
    source(PY+'main.py','events = service.database.get_events',6)
    p(doc,'_event先写run_events，再通过Redis发布“有新事件”的信号。SSE始终从数据库读事实，连接内cursor随着ID前移；Redis故障不会直接删除历史事件。前端timeline监听器调用refreshRun，GET任务详情，读取后端已保存的事件和state快照。它没有直接把SSE的data当成最终页面State。')
    p(doc,'当前eventUrl没有after参数，监听器也没有持久保存event.lastEventId；FastAPI和网关没有读取Last-Event-ID。服务端已具备显式after查询能力，但客户端断线游标续传还未接完整。这个边界需要在面试时说清，不能沿用上一版“客户端已记录游标”的表述。')

    h('在纸上独立画图的验收题',1)
    table(doc,['练习','对照答案'],[
       ['用2分钟画入口','React → Java Security与代理 → FastAPI create_run → TaskDispatcher → Worker → AgentService.start_run → graph.invoke'],
       ['用3分钟画图内拓扑','15节点；四路取证与两路补充；approval两出口；reject跳过verify；没有回环'],
       ['标出State变化','初始4字段；采集写专属字段；evidence累加；stats、hypotheses、plan、approval、results、verification、report依次产生'],
       ['任选一个查询向下跟','Node → Remote adapter → caller.call → client.call_tool → @mcp.tool → tools.py → 参数化SQL；返回反向流动'],
       ['解释暂停和恢复','run_id=thread_id；interrupt暂停；Command(resume)重入approval；数据库approval先于Action写入'],
       ['解释两种终态','业务任务completed只代表流程到END；当前固定三批只有2批连续达标，所以质量observing'],
    ],[4.0,12.8])
    p(doc,'不看答案完成后，再打开graph.py的build逐条核对。如果漏了knowledge_data以外的三个采集节点，或把reject接到verify，说明尚未记住拓扑；如果能画拓扑却说不清get_recent_quality在哪个Server注册，则继续复习七种MCP调用定位。掌握与否以你能否独立画出并解释这些位置为准，文档本身不能代替练习。')
